from __future__ import annotations
import csv, json, os
from pathlib import Path
from typing import Any

def _clean_cell(value: Any) -> Any:
    if value is None: return ""
    if hasattr(value, "isoformat"): return value.isoformat()
    return value

def _rows_to_records(rows: list[list[Any]]) -> tuple[list[str], list[dict[str, Any]]]:
    rows = [[_clean_cell(v) for v in row] for row in rows if any(_clean_cell(v) != "" for v in row)]
    if not rows: return [], []
    headers=[]
    for index,value in enumerate(rows[0]):
        name=str(value).strip() or f"column_{index+1}"; base=name; suffix=2
        while name in headers: name=f"{base}_{suffix}"; suffix+=1
        headers.append(name)
    return headers, [{headers[i]: row[i] if i < len(row) else "" for i in range(len(headers))} for row in rows[1:]]

def _read_tabular(path: Path) -> dict[str, Any]:
    suffix=path.suffix.lower()
    if suffix in {".csv", ".tsv"}:
        delimiter="\t" if suffix==".tsv" else ","
        with path.open("r", encoding="utf-8-sig", newline="") as handle: rows=list(csv.reader(handle, delimiter=delimiter))
        headers,records=_rows_to_records(rows)
        return {"type":"table","sheets":[{"name":path.stem,"headers":headers,"rows":records}]}
    if suffix in {".xlsx", ".xlsm"}:
        from openpyxl import load_workbook
        workbook=load_workbook(path, read_only=True, data_only=True); sheets=[]
        for worksheet in workbook.worksheets:
            rows=[list(row) for row in worksheet.iter_rows(values_only=True)]
            headers,records=_rows_to_records(rows)
            sheets.append({"name":worksheet.title,"headers":headers,"rows":records})
        return {"type":"table","sheets":sheets}
    if suffix == ".xls":
        # 旧版 Excel 使用 xlrd 读取；与 openpyxl 分开处理，避免把 .xls 当作 OOXML。
        try:
            import xlrd
        except ImportError as exc:
            raise ValueError("读取 .xls 需要安装 xlrd==2.0.1，请先执行 pip install xlrd==2.0.1") from exc
        workbook = xlrd.open_workbook(str(path), on_demand=True)
        sheets=[]
        for sheet in workbook.sheets():
            rows=[]
            for row_index in range(sheet.nrows):
                values=[]
                for col_index, value in enumerate(sheet.row_values(row_index)):
                    cell_type = sheet.cell_type(row_index, col_index)
                    if cell_type == xlrd.XL_CELL_DATE:
                        try:
                            value = xlrd.xldate.xldate_as_datetime(value, workbook.datemode).isoformat()
                        except Exception:
                            value = str(value)
                    values.append(value)
                rows.append(values)
            headers,records=_rows_to_records(rows)
            sheets.append({"name":sheet.name,"headers":headers,"rows":records})
        workbook.release_resources()
        return {"type":"table","sheets":sheets}
    raise ValueError(f"不是表格文件: {path.name}")

def extract_document(file_path: str, max_rows: int | None = 200) -> dict[str, Any]:
    path=Path(file_path); suffix=path.suffix.lower()
    if suffix in {".csv", ".tsv", ".xlsx", ".xlsm", ".xls"}:
        result=_read_tabular(path)
        for sheet in result.get("sheets", []): sheet["rows"] = sheet.get("rows", []) if max_rows is None else sheet.get("rows", [])[:max_rows]
        return {"file_name":path.name, **result}
    if suffix==".pdf":
        import pdfplumber
        pages=[]
        with pdfplumber.open(path) as pdf:
            for index,page in enumerate(pdf.pages): pages.append({"page":index+1,"text":page.extract_text() or "","tables":page.extract_tables() or []})
        return {"file_name":path.name,"type":"pdf","pages":pages}
    if suffix==".docx":
        from docx import Document
        document=Document(path); paragraphs=[p.text.strip() for p in document.paragraphs if p.text.strip()]
        tables=[[[cell.text.strip() for cell in row.cells] for row in table.rows] for table in document.tables]
        return {"file_name":path.name,"type":"docx","paragraphs":paragraphs,"tables":tables}
    if suffix in {".txt", ".md", ".log"}: return {"file_name":path.name,"type":"text","text":path.read_text(encoding="utf-8", errors="replace")}
    if suffix==".json": return {"file_name":path.name,"type":"json","data":json.loads(path.read_text(encoding="utf-8"))}
    raise ValueError(f"暂不支持的文件类型: {suffix or '无扩展名'}")

def inspect_spreadsheet(file_path: str, max_rows: int=20) -> dict[str, Any]:
    result=_read_tabular(Path(file_path))
    for sheet in result["sheets"]: sheet["rows"]=sheet["rows"][:max_rows]; sheet["row_count_preview"]=len(sheet["rows"])
    return {"file_name":Path(file_path).name, **result}

def merge_spreadsheets(file_paths: list[str], output_path: str, sheet_name: str="merged") -> dict[str, Any]:
    all_records=[]; headers=[]
    for file_path in file_paths:
        source=Path(file_path); data=_read_tabular(source)
        for sheet in data["sheets"]:
            for record in sheet["rows"]:
                record={str(k):_clean_cell(v) for k,v in record.items()}; record["source_file"]=source.name; record["source_sheet"]=sheet["name"]; all_records.append(record)
                for key in record:
                    if key not in headers: headers.append(key)
    if not headers: raise ValueError("输入表格中没有可合并的数据")
    from openpyxl import Workbook
    from openpyxl.utils import get_column_letter
    workbook=Workbook(); worksheet=workbook.active; worksheet.title=sheet_name[:31] or "merged"; worksheet.append(headers)
    for record in all_records: worksheet.append([record.get(header, "") for header in headers])
    worksheet.freeze_panes="A2"; worksheet.auto_filter.ref=worksheet.dimensions
    for column_index,header in enumerate(headers,1):
        values=[str(header)]+[str(record.get(header,"")) for record in all_records[:100]]; worksheet.column_dimensions[get_column_letter(column_index)].width=min(max(len(value) for value in values)+2,40)
    os.makedirs(Path(output_path).parent, exist_ok=True); workbook.save(output_path)
    return {"file_name":Path(output_path).name,"sheet_name":worksheet.title,"headers":headers,"row_count":len(all_records),"source_count":len(file_paths)}

