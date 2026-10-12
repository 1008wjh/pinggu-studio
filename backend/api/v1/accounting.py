"""会计核算四步流水线 API。

接口形态参考原项目试卷批改：提交后图停在 interrupt，查询待审内容，再使用 Command(resume=...) 恢复。
数据库额外保存每次 State 快照和人工动作，使 MemorySaver 之外仍有可审计记录。
"""
from __future__ import annotations
import csv, io, json, shutil, tempfile
from pathlib import Path
from typing import Annotated, Any, Literal
from uuid import UUID, uuid4
from fastapi import APIRouter, Depends, File, Form, HTTPException, UploadFile
from fastapi.responses import FileResponse
from langgraph.types import Command
from pydantic import BaseModel, Field
from sqlalchemy import text
from sqlalchemy.ext.asyncio import AsyncSession
from backend.agents.accounting.graph import build_accounting_graph
from backend.agents.accounting.state import STEP1, STEP2, STEP3, STEP4
from backend.dependencies import get_current_user, get_db
from backend.services.document_workspace import extract_document

router = APIRouter()
_ROOT = Path(tempfile.gettempdir()) / "pinggu_agent_accounting"
_ROOT.mkdir(parents=True, exist_ok=True)
_GRAPH = build_accounting_graph()


def _uploaded_table_payload(path: str, step: str) -> dict[str, Any]:
    """将人工线下修改后的工作簿转成当前步骤可消费的结构。保留所有工作表。"""
    document = extract_document(path, max_rows=None)
    sheets = document.get("sheets", [])
    if step == STEP1:
        invoices = []
        seen_numbers: set[str] = set()
        for sheet in sheets:
            for index, row in enumerate(sheet.get("rows", []), 1):
                number = str(row.get("数电发票号码") or row.get("发票号码") or "").strip()
                if not number or str(row.get("序号", "")).strip() in {"合计", "合计数"}:
                    continue
                if number in seen_numbers:
                    continue
                seen_numbers.add(number)
                subject = str(row.get("科目名称") or row.get("分类标签") or "待确认/其他")
                amount = float(row.get("金额") or 0); tax = float(row.get("税额") or 0); total = float(row.get("价税合计") or amount + tax)
                invoices.append({"invoice_id": f"INV-UP-{len(invoices)+1:04d}", "source_file": document.get("file_name", "人工修改表格"), "ocr_text": json.dumps(row, ensure_ascii=False, default=str), "label": subject, "matched_rule": "人工上传审核表格", "confidence": 1.0, "invoice_number": number, "invoice_date": str(row.get("开票日期", "")), "seller": str(row.get("销方名称", "")), "buyer": str(row.get("购买方名称", "")), "description": str(row.get("货物或应税劳务名称", "")), "line_count": 1, "amount": amount, "tax_amount": tax, "total_amount": total, "invoice_direction": "sales" if "本单位" in str(row.get("销方名称", "")) else "purchase", "recommended_debit_account": subject, "recommended_credit_account": "主营业务收入"})
        return {"invoices": invoices, "table_preview": [{"workbook": document.get("file_name", ""), "sheet": s.get("name", ""), "headers": s.get("headers", []), "rows": s.get("rows", [])[:8], "row_count": len(s.get("rows", []))} for s in sheets]}
    if step == STEP2:
        vouchers = []
        by_no: dict[str, list[dict[str, Any]]] = {}
        for sheet in sheets:
            for row in sheet.get("rows", []):
                no = str(row.get("凭证号") or "").strip()
                if no and no not in {"凭证号", "合计"}:
                    by_no.setdefault(no, []).append(row)
        for no, rows in by_no.items():
            entries = []
            for row in rows:
                debit = float(row.get("借方金额") or 0); credit = float(row.get("贷方金额") or 0)
                if debit: entries.append({"account": str(row.get("会计科目名称") or ""), "direction": "debit", "amount": debit})
                if credit: entries.append({"account": str(row.get("会计科目名称") or ""), "direction": "credit", "amount": credit})
            total_debit = sum(float(e["amount"]) for e in entries if e["direction"] == "debit")
            total_credit = sum(float(e["amount"]) for e in entries if e["direction"] == "credit")
            vouchers.append({"voucher_no": f"记-{no}", "summary": str(rows[0].get("摘要") or "人工修改凭证"), "debit_account": next((e["account"] for e in entries if e["direction"] == "debit"), ""), "credit_account": next((e["account"] for e in entries if e["direction"] == "credit"), ""), "debit_amount": total_debit, "credit_amount": total_credit, "net_amount": total_debit, "tax_amount": 0, "source_invoice": "", "invoice_number": "", "bank_transaction_id": "", "match_status": "manual", "amount_difference": abs(total_debit-total_credit), "risk": 0.1, "entries": entries})
        return {"vouchers": vouchers, "balanced": all(abs(v["debit_amount"]-v["credit_amount"]) < 0.01 for v in vouchers), "table_preview": [{"workbook": document.get("file_name", ""), "sheet": s.get("name", ""), "headers": s.get("headers", []), "rows": s.get("rows", [])[:8], "row_count": len(s.get("rows", []))} for s in sheets]}
    return {"table_preview": [{"workbook": document.get("file_name", ""), "sheet": s.get("name", ""), "headers": s.get("headers", []), "rows": s.get("rows", [])[:8], "row_count": len(s.get("rows", []))} for s in sheets]}

def _materialize_table_files(state: dict[str, Any], workflow_id: str) -> dict[str, Any]:
    """把每一步的 CSV 工作表组装成可预览、可下载的 xlsx 工作簿。"""
    step = state.get("current_step")
    output = state.get("step_outputs", {}).get(step, {})
    table_csv = output.get("table_csv") or {}
    if not table_csv:
        return state
    folder = _ROOT / workflow_id / "tables"
    folder.mkdir(parents=True, exist_ok=True)
    grouped: dict[str, list[tuple[str, str]]] = {}
    for key, csv_text in table_csv.items():
        workbook_name, separator, sheet_name = str(key).partition("::")
        if not separator:
            workbook_name, sheet_name = "核算结果.xlsx", workbook_name
        grouped.setdefault(workbook_name, []).append((sheet_name or "Sheet1", csv_text))
    from openpyxl import Workbook
    files = []
    report_files = []
    for index, (name, sheets) in enumerate(grouped.items()):
        target = folder / Path(name).name
        # 第三步使用项目样例《报表.xlsx》作为格式模板，保留固定行号、列宽、字体、边框、合计公式和三张表的版式。
        if step == STEP3 and Path(name).name == "三大财务报表.xlsx":
            template = Path(__file__).resolve().parents[3] / "samples" / "例子" / "报表.xlsx"
            if template.exists():
                _materialize_statement_template(template, target, sheets)
                files.append({"key": str(index), "name": target.name, "download_url": f"/api/v1/accounting/workflows/{workflow_id}/tables/{index}"})
                report_folder = _ROOT / workflow_id / "reports"
                report_folder.mkdir(parents=True, exist_ok=True)
                from openpyxl import load_workbook
                for report_key, sheet_name in (("balance_sheet", "资产负债表"), ("income_statement", "利润表"), ("cash_flow_statement", "现金流量表")):
                    report_path = report_folder / f"{sheet_name}.xlsx"
                    report_book = load_workbook(target)
                    for workbook_sheet in list(report_book.worksheets):
                        if workbook_sheet.title != sheet_name:
                            report_book.remove(workbook_sheet)
                    report_book.save(report_path)
                    report_files.append({"key": report_key, "name": report_path.name, "download_url": f"/api/v1/accounting/workflows/{workflow_id}/reports/{report_key}"})
                continue
        workbook = Workbook()
        workbook.remove(workbook.active)
        used_names: set[str] = set()
        for sheet_name, csv_text in sheets:
            base = (sheet_name or "Sheet1")[:31]
            safe = base
            suffix = 2
            while safe in used_names:
                safe = f"{base[:27]}_{suffix}"
                suffix += 1
            used_names.add(safe)
            sheet = workbook.create_sheet(safe)
            rows = list(csv.reader(io.StringIO(csv_text)))
            for row in rows:
                sheet.append(row)
            sheet.freeze_panes = "A2"
            if sheet.max_row and sheet.max_column:
                sheet.auto_filter.ref = sheet.dimensions
        workbook.save(target)
        files.append({"key": str(index), "name": target.name, "download_url": f"/api/v1/accounting/workflows/{workflow_id}/tables/{index}"})
        if step == STEP3:
            report_folder = _ROOT / workflow_id / "reports"
            report_folder.mkdir(parents=True, exist_ok=True)
            for sheet_name, csv_text in sheets:
                report_key = {"资产负债表": "balance_sheet", "利润表": "income_statement", "现金流量表": "cash_flow_statement"}.get(sheet_name)
                if not report_key:
                    continue
                report_path = report_folder / f"{sheet_name}.xlsx"
                report_book = Workbook(); report_sheet = report_book.active; report_sheet.title = sheet_name[:31]
                for row in csv.reader(io.StringIO(csv_text)):
                    report_sheet.append(row)
                report_sheet.freeze_panes = "A2"; report_book.save(report_path)
                report_files.append({"key": report_key, "name": report_path.name, "download_url": f"/api/v1/accounting/workflows/{workflow_id}/reports/{report_key}"})
    updated = dict(output)
    updated["table_files"] = files
    if step == STEP3:
        updated["report_files"] = report_files
    outputs = dict(state.get("step_outputs", {})); outputs[step] = updated; state["step_outputs"] = outputs
    context = dict(state.get("review_context") or {})
    if context.get("step") == step:
        display = dict(context.get("display_data") or {}); display["table_files"] = files; display["report_files"] = report_files if step == STEP3 else display.get("report_files", []); context["display_data"] = display
        ai = dict(context.get("ai_output") or {}); ai["table_files"] = files; ai["report_files"] = report_files if step == STEP3 else ai.get("report_files", []); context["ai_output"] = ai
        state["review_context"] = context
    return state


def _clean_label(value: Any) -> str:
    return "".join(str(value or "").replace("\u3000", "").split())


def _materialize_statement_template(template: Path, target: Path, sheets: list[tuple[str, str]]) -> None:
    """按样例报表.xlsx 填值，保留模板样式和固定结构。"""
    from openpyxl import load_workbook

    workbook = load_workbook(template)
    for sheet_name, csv_text in sheets:
        if sheet_name not in workbook.sheetnames:
            continue
        sheet = workbook[sheet_name]
        rows = list(csv.reader(io.StringIO(csv_text)))
        if not rows:
            continue
        headers = rows[0]
        data = [dict(zip(headers, row)) for row in rows[1:]]
        if sheet_name == "资产负债表":
            value_columns = (4, 5, 8, 9)
            for row_index in range(6, sheet.max_row + 1):
                for column in value_columns:
                    cell = sheet.cell(row_index, column)
                    if not (isinstance(cell.value, str) and cell.value.startswith("=")):
                        cell.value = None
            for item in data:
                left_label = _clean_label(item.get("项目"))
                right_label = _clean_label(item.get("负债和所有者权益"))
                left_no = _number(item.get("序号")); right_no = _number(item.get("负债序号"))
                left_row = next((r for r in range(1, sheet.max_row + 1) if left_no is not None and sheet.cell(r, 3).value == int(left_no)), None)
                right_row = next((r for r in range(1, sheet.max_row + 1) if right_no is not None and sheet.cell(r, 7).value == int(right_no)), None)
                left_row = left_row or next((r for r in range(1, sheet.max_row + 1) if _clean_label(sheet.cell(r, 2).value) == left_label), None)
                right_row = right_row or next((r for r in range(1, sheet.max_row + 1) if _clean_label(sheet.cell(r, 6).value) == right_label), None)
                if left_row:
                    sheet.cell(left_row, 4).value = _number(item.get("期末余额"))
                    sheet.cell(left_row, 5).value = _number(item.get("年初余额"))
                if right_row:
                    sheet.cell(right_row, 8).value = _number(item.get("负债期末余额"))
                    sheet.cell(right_row, 9).value = _number(item.get("负债年初余额"))
        elif sheet_name == "利润表":
            for row_index in range(6, sheet.max_row + 1):
                for column in (4, 5, 6):
                    cell = sheet.cell(row_index, column)
                    if not (isinstance(cell.value, str) and cell.value.startswith("=")):
                        cell.value = None
            for item in data:
                line_no = _number(item.get("行次"))
                row_index = next((r for r in range(1, sheet.max_row + 1) if line_no is not None and sheet.cell(r, 3).value == int(line_no)), None)
                row_index = row_index or next((r for r in range(1, sheet.max_row + 1) if _clean_label(sheet.cell(r, 2).value) == _clean_label(item.get("项目"))), None)
                if row_index:
                    current = _number(item.get("本月数")); cumulative = _number(item.get("本年累计数"))
                    sheet.cell(row_index, 4).value = current; sheet.cell(row_index, 5).value = cumulative; sheet.cell(row_index, 6).value = cumulative
        elif sheet_name == "现金流量表":
            for row_index in range(6, sheet.max_row + 1):
                for column in (4, 5, 6, 7):
                    cell = sheet.cell(row_index, column)
                    if not (isinstance(cell.value, str) and cell.value.startswith("=")):
                        cell.value = None
            for item in data:
                line_no = _number(item.get("行次"))
                row_index = next((r for r in range(1, sheet.max_row + 1) if line_no is not None and sheet.cell(r, 3).value == int(line_no)), None)
                row_index = row_index or next((r for r in range(1, sheet.max_row + 1) if _clean_label(sheet.cell(r, 2).value) == _clean_label(item.get("项目"))), None)
                if row_index:
                    sheet.cell(row_index, 4).value = _number(item.get("本期金额")); sheet.cell(row_index, 5).value = _number(item.get("上年同期金额")); sheet.cell(row_index, 6).value = _number(item.get("本年累计数")); sheet.cell(row_index, 7).value = _number(item.get("本年累计数"))
    workbook.save(target)


def _number(value: Any) -> float | None:
    if value in (None, "", " "):
        return None
    try:
        return float(value)
    except (TypeError, ValueError):
        return None

async def _save(upload: UploadFile | None, folder: Path) -> str:
    if not upload or not upload.filename: return ""
    folder.mkdir(parents=True, exist_ok=True)
    path = folder / Path(upload.filename).name
    with path.open("wb") as handle: shutil.copyfileobj(upload.file, handle)
    return str(path)

def _config(thread_id: str) -> dict[str, Any]:
    return {"configurable": {"thread_id": thread_id}}

def _json(value: Any) -> str:
    return json.dumps(value, ensure_ascii=False, default=str)

async def _snapshot(thread_id: str) -> dict[str, Any]:
    snapshot = await _GRAPH.aget_state(_config(thread_id))
    return dict(snapshot.values or {})

async def _persist(db: AsyncSession, state: dict[str, Any], workflow_id: str, user: dict[str, Any], event: str, action: dict[str, Any] | None = None) -> None:
    state_json = _json(state)
    await db.execute(text("""UPDATE accounting_workflows SET status=:status, current_step=:step,
        state=CAST(:state AS JSONB), checkpoint_version=:version, updated_at=NOW()
        WHERE id=:id AND tenant_id=:tenant_id"""), {"status": state.get("status", "pending_review"), "step": state.get("current_step", STEP1),
        "state": state_json, "version": state.get("checkpoint_version", 0), "id": workflow_id, "tenant_id": user["tenant_id"]})
    await db.execute(text("""INSERT INTO accounting_checkpoints (id, workflow_id, thread_id, step_name, state)
        VALUES (:id, :workflow_id, :thread_id, :step, CAST(:state AS JSONB))"""), {"id": uuid4(), "workflow_id": workflow_id,
        "thread_id": state.get("thread_id", workflow_id), "step": state.get("current_step", STEP1), "state": state_json})
    if action:
        await db.execute(text("""UPDATE accounting_step_reviews SET decision=:decision, edited_output=CAST(:edited AS JSONB),
            comment=:comment, actor_user_id=:actor, acted_at=NOW()
            WHERE id=(SELECT id FROM accounting_step_reviews WHERE workflow_id=:workflow_id AND step_name=:step
                AND decision='pending_review' ORDER BY created_at DESC LIMIT 1)"""), {
            "decision": action.get("action", "return"), "edited": _json(action.get("edited_data") or {}), "comment": action.get("comment", ""),
            "actor": user["user_id"], "workflow_id": workflow_id, "step": action.get("review_step", state.get("current_step", STEP1))})
    context = state.get("review_context") or {}
    if context and state.get("status") == "pending_review":
        step_name = context.get("step", state.get("current_step", STEP1))
        await db.execute(text("""INSERT INTO accounting_step_reviews
            (id, workflow_id, step_name, attempt, ai_output, original_data, confidence, risk, decision)
            VALUES (:id, CAST(:workflow_id AS UUID), CAST(:step_name AS VARCHAR),
                COALESCE((SELECT MAX(attempt)+1 FROM accounting_step_reviews
                    WHERE workflow_id=CAST(:lookup_workflow_id AS UUID) AND step_name=CAST(:lookup_step AS VARCHAR)),1),
                CAST(:output AS JSONB), CAST(:original AS JSONB), :confidence, :risk, 'pending_review')"""), {
            "id": uuid4(), "workflow_id": workflow_id, "step_name": step_name, "lookup_workflow_id": workflow_id, "lookup_step": step_name,
            "output": _json(context.get("ai_output", {})), "original": _json(context.get("original_data", {})),
            "confidence": state.get("confidence", 0), "risk": state.get("risk", 1)})
    await db.execute(text("""INSERT INTO accounting_audit_logs (id, workflow_id, step_name, event_type, payload, actor_user_id)
        VALUES (:id, :workflow_id, :step, :event, CAST(:payload AS JSONB), :actor)"""), {"id": uuid4(), "workflow_id": workflow_id,
        "step": state.get("current_step", STEP1), "event": event, "payload": _json(action or {"status": state.get("status")}), "actor": user["user_id"]})

async def _ensure_memory_checkpoint(row: dict[str, Any]) -> None:
    """进程重启后把 PostgreSQL 最后快照重新挂回 LangGraph 的 human_gate。"""
    config = _config(row["thread_id"])
    snapshot = await _GRAPH.aget_state(config)
    if snapshot.values and snapshot.next:
        return
    stored = row.get("state") or {}
    if isinstance(stored, str):
        stored = json.loads(stored)
    node = {STEP1: "step1", STEP2: "step2", STEP3: "step3", STEP4: "step4"}.get(stored.get("current_step", STEP1), "step1")
    await _GRAPH.aupdate_state(config, stored, as_node=node)
    await _GRAPH.ainvoke(None, config=config)
async def _check_workflow(db: AsyncSession, workflow_id: str, user: dict[str, Any]) -> dict[str, Any]:
    row = (await db.execute(text("SELECT id, thread_id, state, status, current_step, checkpoint_version FROM accounting_workflows WHERE id=:id AND tenant_id=:tenant_id"),
        {"id": workflow_id, "tenant_id": user["tenant_id"]})).mappings().first()
    if not row: raise HTTPException(404, detail="会计核算工作流不存在")
    return dict(row)

class ResumeRequest(BaseModel):
    action: Literal["approve", "edit", "return"]
    edited_data: dict[str, Any] = Field(default_factory=dict)
    comment: str = Field("", max_length=4000)

@router.post("/workflows")
async def create_workflow(
    invoice_files: Annotated[list[UploadFile], File(...)],
    rules_file: UploadFile | None = File(None),
    bank_file: UploadFile | None = File(None),
    history_files: Annotated[list[UploadFile] | None, File()] = None,
    use_llm: bool = Form(True),
    db: AsyncSession = Depends(get_db), user: dict = Depends(get_current_user),
):
    if not invoice_files: raise HTTPException(400, detail="至少上传一份发票图片或 PDF")
    workflow_id = str(uuid4()); folder = _ROOT / workflow_id
    invoice_paths = [await _save(item, folder) for item in invoice_files]
    rules_path = await _save(rules_file, folder); bank_path = await _save(bank_file, folder)
    history_paths = [path for item in (history_files or []) if (path := await _save(item, folder))]
    initial = {"workflow_id": workflow_id, "thread_id": workflow_id, "tenant_id": user["tenant_id"], "user_id": user["user_id"],
        "invoice_paths": invoice_paths, "rules_path": rules_path, "bank_path": bank_path, "history_paths": history_paths,
        "use_llm": use_llm, "current_step": STEP1, "status": "running", "step_outputs": {}, "audit_log": [], "checkpoint_version": 0}
    await db.execute(text("""INSERT INTO accounting_workflows (id, tenant_id, user_id, thread_id, status, current_step, state)
        VALUES (:id, :tenant_id, :user_id, :thread_id, 'running', :step, CAST(:state AS JSONB))"""), {"id": workflow_id, "tenant_id": user["tenant_id"],
        "user_id": user["user_id"], "thread_id": workflow_id, "step": STEP1, "state": _json(initial)})
    try:
        await _GRAPH.ainvoke(initial, config=_config(workflow_id))
        state = await _snapshot(workflow_id)
        state = _materialize_table_files(state, workflow_id)
        await _persist(db, state, workflow_id, user, "workflow_started")
    except Exception as exc:
        raise HTTPException(400, detail={"code": "ACCOUNTING_START_FAILED", "message": str(exc)}) from exc
    return {"workflow_id": workflow_id, "status": state.get("status", "pending_review"), "current_step": state.get("current_step", STEP1),
            "review": state.get("review_context", {}), "checkpoint_version": state.get("checkpoint_version", 0)}

@router.get("/workflows")
async def list_workflows(db: AsyncSession = Depends(get_db), user: dict = Depends(get_current_user)):
    rows=(await db.execute(text("""SELECT id, status, current_step, checkpoint_version, created_at, updated_at
        FROM accounting_workflows WHERE tenant_id=:tenant_id ORDER BY updated_at DESC"""), {"tenant_id": user["tenant_id"]})).mappings().all()
    return {"items":[dict(row) for row in rows]}

@router.delete("/workflows/{workflow_id}")
async def discard_workflow(workflow_id: str, db: AsyncSession = Depends(get_db), user: dict = Depends(get_current_user)):
    """丢弃未完成工作流。页面自然离开不会调用此接口，只有用户明确点击退出才会清理。"""
    row = await _check_workflow(db, workflow_id, user)
    if row["status"] == "completed":
        raise HTTPException(409, detail="已完成工作流不能丢弃")
    await db.execute(text("DELETE FROM accounting_workflows WHERE id=:id AND tenant_id=:tenant_id"), {"id": workflow_id, "tenant_id": user["tenant_id"]})
    folder = _ROOT / workflow_id
    if folder.exists():
        shutil.rmtree(folder, ignore_errors=True)
    return {"deleted": True, "workflow_id": workflow_id}

@router.get("/workflows/{workflow_id}")
async def get_workflow(workflow_id: str, db: AsyncSession = Depends(get_db), user: dict = Depends(get_current_user)):
    row=await _check_workflow(db, workflow_id, user)
    state=await _snapshot(row["thread_id"])
    if not state: state=row.get("state") or {}
    # MemorySaver 的旧快照可能没有 API 生成的下载元数据，按当前状态补齐三个报表文件。
    state = _materialize_table_files(state, workflow_id)
    reviews=(await db.execute(text("""SELECT step_name, attempt, ai_output, original_data, confidence, risk, decision, edited_output, comment, created_at, acted_at
        FROM accounting_step_reviews WHERE workflow_id=:id ORDER BY created_at"""), {"id":workflow_id})).mappings().all()
    audit=(await db.execute(text("SELECT step_name, event_type, payload, actor_user_id, created_at FROM accounting_audit_logs WHERE workflow_id=:id ORDER BY created_at"), {"id":workflow_id})).mappings().all()
    return {"workflow_id":workflow_id,"status":state.get("status",row["status"]),"current_step":state.get("current_step",row["current_step"]),
        "confidence":state.get("confidence",0),"risk":state.get("risk",1),"review":state.get("review_context",{}),
        "step_outputs":state.get("step_outputs",{}),"reviews":[dict(item) for item in reviews],"audit_logs":[dict(item) for item in audit],
        "checkpoint_version":state.get("checkpoint_version",row["checkpoint_version"])}

@router.post("/workflows/{workflow_id}/resume")
async def resume_workflow(workflow_id: str, req: ResumeRequest, db: AsyncSession = Depends(get_db), user: dict = Depends(get_current_user)):
    row=await _check_workflow(db, workflow_id, user)
    if row["status"] == "completed": raise HTTPException(409, detail="工作流已经完成")
    previous_state = await _snapshot(row["thread_id"])
    previous_step = previous_state.get("current_step", row.get("current_step", STEP1))
    decision={"action":req.action,"edited_data":req.edited_data,"comment":req.comment,"user_id":user["user_id"],"review_step":previous_step}
    try:
        await _ensure_memory_checkpoint(row)
        await _GRAPH.ainvoke(Command(resume=decision), config=_config(row["thread_id"]))
        state=await _snapshot(row["thread_id"])
        state = _materialize_table_files(state, workflow_id)
        await _persist(db,state,workflow_id,user,f"human_{req.action}",decision)
    except Exception as exc:
        raise HTTPException(400, detail={"code":"ACCOUNTING_RESUME_FAILED","message":str(exc)}) from exc
    return {"workflow_id":workflow_id,"status":state.get("status"),"current_step":state.get("current_step"),"review":state.get("review_context",{}),
        "step_outputs":state.get("step_outputs",{}),"checkpoint_version":state.get("checkpoint_version",0)}

@router.get("/workflows/{workflow_id}/reports/{report_key}")
async def download_statement_report(workflow_id: str, report_key: Literal["balance_sheet", "income_statement", "cash_flow_statement"], db: AsyncSession = Depends(get_db), user: dict = Depends(get_current_user)):
    row = await _check_workflow(db, workflow_id, user)
    path = _ROOT / workflow_id / "reports" / {"balance_sheet": "资产负债表.xlsx", "income_statement": "利润表.xlsx", "cash_flow_statement": "现金流量表.xlsx"}[report_key]
    if not path.exists():
        state = row.get("state") or {}
        if isinstance(state, str): state = json.loads(state)
        statement_state = dict(state)
        statement_state["current_step"] = STEP3
        state = _materialize_table_files(statement_state, workflow_id)
    if not path.exists():
        raise HTTPException(404, detail="该工作流尚未生成三大报表")
    return FileResponse(path, filename=path.name, media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")


@router.get("/workflows/{workflow_id}/tables/{file_key}")
async def download_accounting_table(workflow_id: str, file_key: str, db: AsyncSession = Depends(get_db), user: dict = Depends(get_current_user)):
    row = await _check_workflow(db, workflow_id, user)
    folder = _ROOT / workflow_id / "tables"
    state = await _snapshot(row["thread_id"]) or row.get("state") or {}
    if isinstance(state, str):
        state = json.loads(state)
    # MemorySaver 中的快照可能早于 API 物化步骤，先根据当前步骤的 CSV 恢复 xlsx 和 table_files 元数据。
    state = _materialize_table_files(state, workflow_id)
    outputs = state.get("step_outputs", {})
    metadata = None
    for step_name in (STEP1, STEP2, STEP3):
        output = outputs.get(step_name, {})
        metadata = next((item for item in output.get("table_files", []) if str(item.get("key")) == str(file_key)), None)
        if metadata:
            break
    path = folder / Path(metadata.get("name", "")).name if metadata else None
    if (not path or not path.exists()) and outputs:
        for step_name in (STEP1, STEP2, STEP3):
            if outputs.get(step_name, {}).get("table_csv"):
                candidate_state = dict(state)
                candidate_state["current_step"] = step_name
                candidate_state = _materialize_table_files(candidate_state, workflow_id)
                output = candidate_state.get("step_outputs", {}).get(step_name, {})
                metadata = next((item for item in output.get("table_files", []) if str(item.get("key")) == str(file_key)), None)
                if metadata:
                    path = folder / Path(metadata.get("name", "")).name
                    if path.exists():
                        break
    if not path or not path.exists():
        raise HTTPException(404, detail="表格文件不存在")
    return FileResponse(path, filename=path.name, media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")


@router.post("/workflows/{workflow_id}/table-upload")
async def upload_review_table(workflow_id: str, table_file: UploadFile = File(...), db: AsyncSession = Depends(get_db), user: dict = Depends(get_current_user)):
    """用户下载工作簿线下修改后上传，作为当前步骤的正式输入。"""
    row = await _check_workflow(db, workflow_id, user)
    state = await _snapshot(row["thread_id"]) or row.get("state") or {}
    if isinstance(state, str):
        state = json.loads(state)
    step = state.get("current_step", STEP1)
    folder = _ROOT / workflow_id / "uploads"
    uploaded = await _save(table_file, folder)
    if not uploaded or Path(uploaded).suffix.lower() not in {".xlsx", ".xls", ".xlsm", ".csv"}:
        raise HTTPException(400, detail="人工修改结果必须上传 xlsx、xls、xlsm 或 csv 文件")
    payload = _uploaded_table_payload(uploaded, step)
    decision = {"action": "upload", "edited_data": payload, "comment": "用户上传线下修改后的表格", "user_id": user["user_id"], "review_step": step}
    try:
        await _ensure_memory_checkpoint(row)
        await _GRAPH.ainvoke(Command(resume=decision), config=_config(row["thread_id"]))
        state = await _snapshot(row["thread_id"])
        state = _materialize_table_files(state, workflow_id)
        await _persist(db, state, workflow_id, user, "human_upload", decision)
    except Exception as exc:
        raise HTTPException(400, detail={"code": "ACCOUNTING_TABLE_UPLOAD_FAILED", "message": str(exc)}) from exc
    return {"workflow_id": workflow_id, "status": state.get("status"), "current_step": state.get("current_step"), "review": state.get("review_context", {}), "step_outputs": state.get("step_outputs", {}), "checkpoint_version": state.get("checkpoint_version", 0)}

@router.get("/health")
async def accounting_health():
    return {"status":"ok","agent":"accounting_pipeline","steps":[STEP1,STEP2,STEP3,STEP4],"interrupt_required":True,"checkpoint":"memory_plus_postgres"}





